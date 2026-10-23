import re, zipfile
from docx import Document
from docx.oxml.ns import qn
import sys
f = sys.argv[1]
d = Document(f); x = zipfile.ZipFile(f).read('word/document.xml').decode('utf-8')
print('styles:', set(p.style.name for p in d.paragraphs))          # 只应有 Normal
print('numPr:', len(re.findall(r'<w:numPr>', x)))                  # 必须 0
print('List styles:', len(re.findall(r'ListNumber|ListBullet', x)))# 必须 0
print('superscript:', len(re.findall(r'w:vertAlign w:val="superscript"', x)))
names = set(re.findall(r'<w:bookmarkStart [^>]*w:name="([^"]+)"', x))
anchors = set(re.findall(r'<w:hyperlink w:anchor="([^"]+)"', x))
print('unresolved anchors:', [a for a in anchors if a not in names])  # 必须 []

# ★ R27 内部交叉引用：[[§x.y 表]] 必须落到 _Sec… 书签；未解析项会渲染成红字「§…?」
_xref = re.findall(r'<w:hyperlink w:anchor="(_Sec[^"]+)"', x)
print('内部交叉引用 xref: %d 处 -> %d 个目标（失效 %d）'
      % (len(_xref), len(set(_xref)), len([a for a in _xref if a not in names])))
# 未解析的 xref 会渲染成「§…?」——标签里可能有中文（如「附录B 表B1」），故不能只匹配数字
print('xref 未解析残留（底稿中的 §…?）:',
      len(re.findall(r'§[^<]{1,28}?\?', x)))

# ★ R15 附录正文单倍行距（标题除外）＋ pPr 子元素次序合法
SEQ = ["w:pStyle","w:keepNext","w:keepLines","w:pageBreakBefore","w:framePr",
 "w:widowControl","w:numPr","w:suppressLineNumbers","w:pBdr","w:shd","w:tabs",
 "w:suppressAutoHyphens","w:kinsoku","w:wordWrap","w:overflowPunct","w:topLinePunct",
 "w:autoSpaceDE","w:autoSpaceDN","w:bidi","w:adjustRightInd","w:snapToGrid","w:spacing",
 "w:ind","w:contextualSpacing","w:mirrorIndents","w:suppressOverlap","w:jc",
 "w:textDirection","w:textAlignment","w:textboxTightWrap","w:outlineLvl","w:divId",
 "w:cnfStyle","w:rPr","w:sectPr","w:pPrChange"]
_id = {t: i for i, t in enumerate(SEQ)}
bad_order = 0
for pp in re.findall(r'<w:pPr>(.*?)</w:pPr>', x, re.S):
    pos = [_id.get(t, 999) for t in re.findall(r'<(w:[A-Za-z]+)[ />]', pp)]
    bad_order += (pos != sorted(pos))
print('pPr 乱序段落数:', bad_order)                                  # 必须 0

# ★ rPr 子元素次序合法（角标超链接的 rPr 曾把 vertAlign 排在 color/sz 之前）
RSEQ = ["w:rStyle","w:rFonts","w:b","w:bCs","w:i","w:iCs","w:caps","w:smallCaps",
 "w:strike","w:dstrike","w:outline","w:shadow","w:emboss","w:imprint","w:noProof",
 "w:snapToGrid","w:vanish","w:webHidden","w:color","w:spacing","w:w","w:kern",
 "w:position","w:sz","w:szCs","w:highlight","w:u","w:effect","w:bdr","w:shd",
 "w:fitText","w:vertAlign","w:rtl","w:cs","w:em","w:lang","w:eastAsianLayout",
 "w:specVanish","w:oMath"]
_rid = {t: i for i, t in enumerate(RSEQ)}
bad_rorder = 0
for rp in re.findall(r'<w:rPr>(.*?)</w:rPr>', x, re.S):
    pos = [_rid.get(t, 999) for t in re.findall(r'<(w:[A-Za-z]+)[ />]', rp)]
    bad_rorder += (pos != sorted(pos))
print('rPr 乱序 run 数:', bad_rorder)                                 # 必须 0
in_app, ok_single, bad_single, head_wrong = False, 0, 0, 0
for p in re.findall(r'<w:p[ >].*?</w:p>', x.split('<w:body>', 1)[1], re.S):
    txt = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', p))
    head = '<w:outlineLvl' in p
    if head and re.match(r'^附\s*[錄录]\s*[A-Z]', txt):
        in_app = True
    if not in_app:
        continue
    lv = re.search(r'w:line="(\d+)"', p)
    lv = lv.group(1) if lv else None
    if head:
        head_wrong += (lv == '240')          # 标题不得被压成附录正文的单倍（版式档标题行距 1.2/1.5 均合规）
    elif lv == '240':
        ok_single += 1                       # 正文须单倍
    else:
        bad_single += 1
print('附录正文单倍:', ok_single, '| 非单倍:', bad_single,
      '| 标题被误压:', head_wrong)                                   # 期望 0 / 0

# ★ R15-4 文件名日期 = 生成当天 CST(UTC+8)
import datetime, os
today = datetime.datetime.now(
    datetime.timezone(datetime.timedelta(hours=8))).strftime('%Y%m%d')
print('文件名含当天 CST 日期:', today in os.path.basename(f), today)   # 必须 True

# ★ R14/R16 终态自检：过程痕迹 / 稻草人词表，命中必须为 0
# 注：`第一轮/第二轮/首轮/本轮` 有正当用法（"本轮国家战略"），不列入机器否决表，由人工判断（R14-1）
VETO = ['回填','数据版本','较首次','此前被','应修正为','待补证','负面发现',
        '并不成立','推断不成立','并非没有','走访前','走访后','补充检索',
        '增量判断','（新增）','本次修订','待核']
cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
hits = {w: sum(w in p.text for p in d.paragraphs) + sum(w in c for c in cells)
        for w in VETO}
print('过程/稻草人词命中:', {k: v for k, v in hits.items() if v})       # 期望 {}

# ★ 多编号引用：必须逐编号成链
for p in d.paragraphs:
    m = re.search(r'\[(\d+(?:,\d+)+)\]', p.text)
    if m:
        exp = m.group(1).split(',')
        got = [h.get(qn('w:anchor')) for h in p._element.findall(qn('w:hyperlink'))]
        ok = got[:len(exp)] == ['_Ref%s' % n for n in exp]
        print('多编号样例', m.group(0), '->', got, 'OK' if ok else 'FAIL')
        break

# ★ 导航窗格标题树：统计 outlineLvl 分布（1=章 2=节 3=子节/附录分组）
lv = {}
for p in d.paragraphs:
    pPr = p._element.find(qn('w:pPr'))
    ol = pPr.find(qn('w:outlineLvl')) if pPr is not None else None
    if ol is not None:
        k = ol.get(qn('w:val')); lv[k] = lv.get(k, 0) + 1
print('outlineLvl:', lv)          # 附录分组应计入 3

# ★ 表格数据引用：统计含引用链接的单元格数与表格数
cells = sum(1 for t in d.tables for r in t.rows for c in r.cells
            if 'w:anchor="_Ref' in c._tc.xml)
print('tables:', len(d.tables), '| cells with citation:', cells)   # 重要表应逐行有

# ★ 僵尸文献：被引用编号集合必须覆盖 1..N 无缺口
nums = sorted({int(m.group(1)) for a in anchors if (m := re.match(r'_Ref(\d+)$', a))})
N = max(nums) if nums else 0
print('cited refs:', len(nums), '| 未被引用的条目:',
      [n for n in range(1, N + 1) if n not in nums])              # 必须 []

# ★ R6.1 推算三件套：标了【推算】却没有公式/参数的段落
bad = []
for p in d.paragraphs:
    t = p.text
    if '【推算】' in t and not re.search(r'[=≈×/*]', t):
        bad.append(t[:60])
print('【推算】缺公式/参数的段落:', len(bad), bad[:3])              # 期望 0
print('【推算】总数:', sum('【推算】' in p.text for p in d.paragraphs))

# ★ R6.3 护栏线索：含"算术和/峰值/合计"且同段有量级数的结论段，人工复核
peak = [p.text[:60] for p in d.paragraphs
        if re.search(r'算术和|峰值相加|相加得出', p.text)]
print('需人工复核"算术和当瞬时负载"的段落:', len(peak), peak[:3])

# ★ R8 问题—建议可追溯：附录是否存在映射表
keys = lambda t: {c.text.strip() for c in t.rows[0].cells}
_f = lambda ks, ws: any(w in k for k in ks for w in ws)
print('问题—建议映射表存在:', any(
    _f(keys(t), ('短板', '问题', '問題', '现状问题', '現狀問題'))
    and _f(keys(t), ('建议', '建議', '方案', '对策', '對策'))
    for t in d.tables))                                            # 期望 True

# ★ R9-5 报告内必须有「当前算力分项汇总表」：含高校行 + 口径标注
hit = None
for t2 in d.tables:
    txt = ' '.join(c.text for r in t2.rows for c in r.cells)
    if ('高校' in txt or '大学' in txt) and ('装机' in txt or '在用' in txt) \
            and ('企业' in txt):
        hit = t2
print('当前算力分项汇总表存在:', hit is not None)                       # 期望 True

# ★ R9-1/2/4 由 parse_survey.py 的台账机器校验（必答题作答率 / 0 值单列 / q4↔q6 比值）
import os
led = os.path.join(os.path.dirname(os.path.abspath(f)), 'survey_current_compute.md')
print('当前算力台账存在:', os.path.exists(led), led)                    # 期望 True
if os.path.exists(led):
    s = open(led, encoding='utf-8').read()
    print('  必答题作答率 100%:', '100%' in s)
    print('  0 值单列:', '0＝无在用算力' in s or '0（0＝无' in s)
    print('  存在越界/未识别条目（需人工处置）:', s.count('⚠️') > 0)

# ★ R10 算力指标"逐值"精确引用：数值单元格必须自带角标（表头行不计）
# ⚠️ 必须与 assets/move_citations_to_value.py 的 MET 完全一致，否则"100%"不可复现
_NUM = r'\d[\d,\.]*\s*[萬万亿億]?\s*'
_UNIT = (r'(?:亿亿次|万亿次|PFLOPS|PFlops|P@FP16|Eops|EFLOPS|EPLOPS'
         r'|卡時|卡时|張|张|卡|台|柜|櫃|机架|機架|机位|機位'
         r'|MW|kW|MVA|kVA|GWh|%|P(?![A-Za-z@]))')
MET = re.compile(_NUM + r'(?:(?:[A-Za-z]{1,6}\s*)?\d*\s*等效\s*)?' + _UNIT)
tt = ss = 0
for t2 in d.tables:
    h0 = t2.rows[0].cells[0]
    # 关键数字带：标题行带角标＝整块卡片的共同来源，卡片内只标额外来源
    band_cited = (h0.text.strip()[:6] in ('本节关键数字', '本節關鍵數字')
                  and 'w:anchor="_Ref' in h0._tc.xml)
    for r2 in t2.rows[1:]:                      # 跳过表头行
        for c2 in r2.cells:
            if MET.search(c2.text):
                tt += 1
                if band_cited or 'w:anchor="_Ref' in c2._tc.xml:
                    ss += 1
print('R10 数值单元格自带角标: %d/%d = %.0f%%' % (ss, tt, 100.0 * ss / max(tt, 1)))  # 目标 100%

# ★ R10-2 引用条目是否"引到条"（含日期/文号/URL/记录级定位之一）
# 注意：转换器把 `Rn.` 条目渲染成 `[n] …`，所以按 `^\[\d+\]` 匹配，不是 `^Rn\.`
weak = []
for p2 in d.paragraphs:
    if re.match(r'^\[\d+\]\s', p2.text.strip()):
        s2 = p2.text
        if not re.search(r'20\d\d|第\s*\d+\s*[条頁页]|URL|https?://|记录\s*#'
                          r'|落盘|\.md|\.csv|\.json|\.pdf|/', s2):
            weak.append(s2[:50])
print('R10 定位信息偏弱的参考文献条目:', len(weak), weak[:3])          # 期望 0

# ★ R7 繁体版红线：文件路径与 URL 未被转换、无台湾用语
PROF_IS_GOV = '標楷體' in x          # gov 档以標楷體为唯一中文字体，可直接判档
if PROF_IS_GOV:
    paths = re.findall(r'`[^`]*\.(?:md|csv|json|pdf|png|xlsx|docx)`', x)
    # 专名白名单：香港官方名册全称「資料中心服務供應商名冊」属登记原文用字，
    # 依 R7 必须保留，不视为台湾用语误用
    x2 = x.replace('資料中心服務供應商名冊', '')
    tw = [w for w in ('軟體', '網路', '資料中心', '晶片', '影片', '平臺', '伺服器') if w in x2]
    print('R7 繁体版：路径/URL 出现 %d 处（应原样保留）；台湾用语命中 %s（应为 []）' % (len(paths), tw))
    print('R7 繁体版：全篇字号是否统一为 14pt（sz=28 为主）:',
          len(re.findall(r'<w:sz w:val="28"', x)) > len(re.findall(r'<w:sz w:val="22"', x)))

# ★ R6.5 官方口径并录：出现"口径""差异"说明的表格或段落
_o = [p2.text for p2 in d.paragraphs if re.search(r'口徑|口径', p2.text) and re.search(r'差異|差异|並錄|并录|待核', p2.text)]
print('R6.5 口径差异说明段数:', len(_o))          # 期望 ≥1（有官方口径冲突时）

# ★ R14 端到端最终态：过程痕迹自检（命中应为 0；"本轮国家战略"等正当用法需人工判断）
PROC = ['第一輪', '第二輪', '第三輪', '首輪', '本輪', '回填', '數據版本', '較首次',
        '此前被系統性低估', '應修正為', '待核', '按R6', '按R10', '按R12']
hits = {k: x.count(k) for k in PROC if x.count(k)}
print('R14 過程痕跡命中:', hits or '無 ✅', '（"本輪國家戰略""雙輪驅動"等屬正當用法）')

# 关键判断段落引用覆盖率
key = re.compile(r'结论|判断|建议|定位|总体架构')
tot = wit = 0
for p in d.paragraphs:
    t = p.text.strip()
    if len(t) < 25 or t.startswith('['): continue
    if key.search(t):
        tot += 1
        if p._element.findall(qn('w:hyperlink')): wit += 1
print('key paragraphs cited: %d/%d' % (wit, tot))                  # 目标 100%